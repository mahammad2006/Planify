import os
from fastapi import FastAPI, Depends, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from database import engine, get_db
from sqlalchemy.orm import Session
import models
import io
import time
import openpyxl
from routers.groups import router as groups_router
from routers.teachers import router as teachers_router
from routers.rooms import router as rooms_router
from routers.subjects import router as subjects_router
from scheduler import generate_schedule

models.Base.metadata.create_all(bind=engine)

app = FastAPI(title="Planify - Universitet Cədvəl Sistemi")

# CORS — FRONTEND_URL env variable varsa onu da əlavə et
allowed_origins = [
    "http://localhost:3000",
    "http://localhost:5173",
]

frontend_url = os.environ.get("FRONTEND_URL")
if frontend_url:
    allowed_origins.append(frontend_url)

app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.include_router(groups_router)
app.include_router(teachers_router)
app.include_router(rooms_router)
app.include_router(subjects_router)

@app.get("/")
def root():
    return {"message": "Planify API işləyir ✅"}

@app.post("/generate-schedule")
def generate(db: Session = Depends(get_db)):
    return generate_schedule(db)

@app.get("/schedule")
def get_schedule(db: Session = Depends(get_db)):
    schedules = db.query(models.Schedule).all()
    day_names = ["Bazar ertəsi", "Çərşənbə axşamı", "Çərşənbə", "Cümə axşamı", "Cümə"]
    result = []
    for s in schedules:
        result.append({
            "id": s.id,
            "subject": s.subject.name,
            "teacher": s.subject.teacher.full_name,
            "group": s.subject.group.name,
            "room": s.room.name,
            "day": day_names[s.day],
            "slot": s.time_slot + 1,
            "shift": s.shift
        })
    return result

@app.post("/import-excel")
async def import_excel(file: UploadFile = File(...), db: Session = Depends(get_db)):
    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents))

    results = {"groups": 0, "teachers": 0, "rooms": 0, "subjects": 0, "errors": []}

    # ── Qruplar ──
    if "Qruplar" in wb.sheetnames:
        ws = wb["Qruplar"]
        for row in list(ws.iter_rows(min_row=2, values_only=True)):
            if not row[0]:
                continue
            try:
                shift = "morning" if str(row[2]).lower() in ["səhər", "morning"] else "afternoon"
                free_day = None
                if row[3] and str(row[3]).strip():
                    day_map = {
                        "bazar ertəsi": 0, "çərşənbə axşamı": 1,
                        "çərşənbə": 2, "cümə axşamı": 3, "cümə": 4,
                        "yoxdur": -1, "avtomatik": None
                    }
                    free_day = day_map.get(str(row[3]).strip().lower(), None)

                existing = db.query(models.Group).filter(models.Group.name == str(row[0])).first()
                if not existing:
                    db.add(models.Group(
                        name=str(row[0]),
                        student_count=int(row[1]),
                        shift=shift,
                        free_day=free_day
                    ))
                    results["groups"] += 1
            except Exception as e:
                results["errors"].append(f"Qrup xətası: {row[0]} — {str(e)}")
        db.commit()

    # ── Müəllimlər ──
    if "Müəllimlər" in wb.sheetnames:
        ws = wb["Müəllimlər"]
        for row in list(ws.iter_rows(min_row=2, values_only=True)):
            if not row[0]:
                continue
            try:
                email = str(row[1]).strip() if row[1] else None
                existing = db.query(models.Teacher).filter(models.Teacher.full_name == str(row[0])).first()
                if not existing:
                    db.add(models.Teacher(
                        full_name=str(row[0]),
                        email=email if email else None
                    ))
                    results["teachers"] += 1
            except Exception as e:
                results["errors"].append(f"Müəllim xətası: {row[0]} — {str(e)}")
        db.commit()

    # ── Otaqlar ──
    if "Otaqlar" in wb.sheetnames:
        ws = wb["Otaqlar"]
        for row in list(ws.iter_rows(min_row=2, values_only=True)):
            if not row[0]:
                continue
            try:
                room_type_map = {
                    "kompüter": "computer", "computer": "computer",
                    "laboratoriya": "laboratory", "laboratory": "laboratory",
                    "adi": "regular", "regular": "regular"
                }
                room_type = room_type_map.get(str(row[2]).strip().lower() if row[2] else "adi", "regular")
                existing = db.query(models.Room).filter(models.Room.name == str(row[0])).first()
                if not existing:
                    db.add(models.Room(
                        name=str(row[0]),
                        capacity=int(row[1]),
                        room_type=room_type
                    ))
                    results["rooms"] += 1
            except Exception as e:
                results["errors"].append(f"Otaq xətası: {row[0]} — {str(e)}")
        db.commit()

    # ── Fənlər ──
    if "Fənlər" in wb.sheetnames:
        ws = wb["Fənlər"]
        for row in list(ws.iter_rows(min_row=2, values_only=True)):
            if not row[0]:
                continue
            try:
                teacher = db.query(models.Teacher).filter(models.Teacher.full_name == str(row[1])).first()
                if not teacher:
                    results["errors"].append(f"Fənn xətası: '{row[1]}' müəllim tapılmadı")
                    continue

                subject_type = "normal"
                if row[6]:
                    type_map = {
                        "birləşmiş": "merged", "merged": "merged",
                        "bölünmüş": "split", "split": "split"
                    }
                    subject_type = type_map.get(str(row[6]).strip().lower(), "normal")

                group_names = [g.strip() for g in str(row[2]).split(",")]

                merge_id = None
                if subject_type == "merged":
                    merge_id = f"merge_{row[0]}_{row[1]}_{int(time.time())}"

                for gname in group_names:
                    group = db.query(models.Group).filter(models.Group.name == gname).first()
                    if not group:
                        results["errors"].append(f"Fənn xətası: '{gname}' qrup tapılmadı")
                        continue

                    requires_computer = str(row[4]).lower() in ["bəli", "yes", "true", "1"] if row[4] else False
                    requires_laboratory = str(row[5]).lower() in ["bəli", "yes", "true", "1"] if row[5] else False
                    hours = int(row[3]) if row[3] else 2

                    if subject_type == "split":
                        counts_str = str(row[7]).strip() if row[7] else ""
                        counts = counts_str.split("+") if "+" in counts_str else []
                        if len(counts) < 2:
                            results["errors"].append(
                                f"Fənn xətası: '{row[0]}' — bölünmə sayları düzgün deyil (məs: 13+12)"
                            )
                            continue

                        split_id = f"split_{row[0]}_{gname}_{int(time.time())}"
                        for cnt in counts:
                            db.add(models.Subject(
                                name=str(row[0]),
                                teacher_id=teacher.id,
                                group_id=group.id,
                                hours_per_week=hours,
                                requires_computer=requires_computer,
                                requires_laboratory=requires_laboratory,
                                subject_type=subject_type,
                                split_id=split_id,
                                split_student_count=int(cnt.strip())
                            ))
                        results["subjects"] += 1
                    else:
                        db.add(models.Subject(
                            name=str(row[0]),
                            teacher_id=teacher.id,
                            group_id=group.id,
                            hours_per_week=hours,
                            requires_computer=requires_computer,
                            requires_laboratory=requires_laboratory,
                            subject_type=subject_type,
                            merge_id=merge_id,
                        ))
                        results["subjects"] += 1

            except Exception as e:
                results["errors"].append(f"Fənn xətası: {row[0]} — {str(e)}")
        db.commit()

    return results
